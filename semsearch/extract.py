"""Text extraction for each supported format. Any failure returns "" so indexing never stops."""
import json
import re
import subprocess
from pathlib import Path

from . import config


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


def _read_text(path: Path, limit=None) -> str:
    raw = path.open("rb").read(limit or config.MAX_TEXT_BYTES)
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="ignore")  # UTF-16 only with a BOM
    if b"\x00" in raw[:4096]:
        return ""  # actually a binary file
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError as e:
        # cut in the middle of a multi-byte character, or a few bad bytes: still UTF-8
        if e.start >= len(raw) - 3 or raw.decode("utf-8", "replace").count("\ufffd") < len(raw) / 1000:
            return raw.decode("utf-8", errors="ignore")
    for enc in ("big5", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _pdf(path: Path) -> str:
    import pypdfium2 as pdfium
    out, total = [], 0
    doc = pdfium.PdfDocument(str(path))
    try:
        for page in doc:
            t = page.get_textpage().get_text_range()
            out.append(t)
            total += len(t)
            if total > config.MAX_CHARS_PER_FILE:
                break
    finally:
        doc.close()
    return "\n".join(out)


def _docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        parts.append(f"[slide {i}] " + "\n".join(t for t in texts if t))
    return "\n".join(parts)


def _xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts, total = [], 0
    for ws in wb.worksheets:
        parts.append(f"[sheet {ws.title}]")
        for row in ws.iter_rows(max_row=300, values_only=True):
            line = " | ".join(str(v) for v in row if v is not None)
            if line:
                parts.append(line)
                total += len(line)
            if total > config.MAX_CHARS_PER_FILE:
                break
    wb.close()
    return "\n".join(parts)


def _ipynb(path: Path) -> str:
    nb = json.loads(_read_text(path, limit=50_000_000) or "{}")  # never truncate: the JSON must stay valid
    parts = []
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        parts.append("".join(src) if isinstance(src, list) else src)
    return "\n".join(parts)


def _textutil(path: Path) -> str:
    # built into macOS; reads .doc / .rtf / .odt
    r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)],
                       capture_output=True, timeout=30)
    return r.stdout.decode("utf-8", errors="ignore")


def _html(path: Path) -> str:
    t = _read_text(path)
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", t)
    return re.sub(r"<[^>]+>", " ", t)


HANDLERS = {
    ".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx,
    ".ipynb": _ipynb, ".doc": _textutil, ".rtf": _textutil, ".odt": _textutil,
    ".html": _html, ".htm": _html,
}


def garbled(text: str) -> bool:
    """PDFs with broken font encodings yield text like "òðòö@ ðò@"; too many Latin-1 letters means garbage."""
    letters = [c for c in text if not c.isspace()]
    if len(letters) < 20:
        return False
    odd = sum(1 for c in letters if "\u00c0" <= c <= "\u00ff")
    return odd / len(letters) > 0.2


def extract(path: Path, ext: str, size: int) -> str:
    if ext not in config.CONTENT_EXT:
        return ""
    if ext in config.DOC_EXT and size > config.MAX_DOC_BYTES:
        return ""
    try:
        text = HANDLERS.get(ext, _read_text)(path)
    except Exception:
        return ""
    text = _clean(text)[: config.MAX_CHARS_PER_FILE]
    return "" if garbled(text) else text


def chunk(text: str):
    """Split into overlapping chunks, preferring to cut at line breaks or sentence ends."""
    n, size, ov = len(text), config.CHUNK_CHARS, config.CHUNK_OVERLAP
    chunks, start = [], 0
    while start < n and len(chunks) < config.MAX_CHUNKS_PER_FILE:
        end = min(n, start + size)
        if end < n:
            window = text[start + size // 2:end]
            cut = max(window.rfind("\n"), window.rfind("。"), window.rfind(". "))
            if cut > 0:
                end = start + size // 2 + cut + 1
        piece = text[start:end].strip()
        if len(piece) >= 20:
            chunks.append(piece)
        if end >= n:
            break
        start = max(end - ov, start + 1)
    return chunks
